import os
import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SRC_TEMPLATE = r"C:\Users\kunal\Desktop\NIRIKSHAK\SIH ppt - kunal\SIH2026-IDEA-Presentation-Format.pptx"
OUT_PPTX = r"C:\Users\kunal\Desktop\NIRIKSHAK\SIH ppt - kunal\SIH26034-KUNAL-PPTv2.pptx"
BASE_DIR = r"C:\Users\kunal\Desktop\NIRIKSHAK\SIH ppt - kunal"
ASSETS_DIR = os.path.join(BASE_DIR, "assests")
DIAGRAMS_DIR = os.path.join(ASSETS_DIR, "diagrams")
UI_SCREENS_DIR = os.path.join(ASSETS_DIR, "ui-screens")
EVIDENCE_DIR = os.path.join(ASSETS_DIR, "05_product_inspection_evidence")
LOGOS_DIR = os.path.join(ASSETS_DIR, "03_brand_and_logos")

# Color definitions matching winning deck (ideal_1..6)
NAVY_TITLE   = RGBColor(0x0F, 0x17, 0x2A) # #0F172A
BLUE_ACCENT  = RGBColor(0x02, 0x84, 0xC7) # #0284C7
BLUE_DARK    = RGBColor(0x03, 0x69, 0xA1) # #0369A1
GOLD_AMBER   = RGBColor(0xD9, 0x77, 0x06) # #D97706
GOLD_LIGHT   = RGBColor(0xFE, 0xF3, 0xC7) # #FEF3C7
EMERALD_ACC  = RGBColor(0x05, 0x96, 0x69) # #059669
GREEN_LIGHT  = RGBColor(0xDC, 0xFC, 0xE7) # #DCFCE7
RED_ACCENT   = RGBColor(0xDC, 0x26, 0x26) # #DC2626
PINK_LIGHT   = RGBColor(0xFE, 0xE2, 0xE2) # #FEE2E2
PURPLE_ACC   = RGBColor(0x7C, 0x3A, 0xED) # #7C3AED
PURPLE_LIGHT = RGBColor(0xF3, 0xE8, 0xFF) # #F3E8FF
CARD_BG_GRAY = RGBColor(0xF8, 0xFA, 0xFC) # #F8FAFC
CARD_BORDER  = RGBColor(0xEB, 0xF0, 0xF5) # #EBF0F5
TEXT_MUTED   = RGBColor(0x47, 0x55, 0x69) # #475569
TEXT_DARK    = RGBColor(0x1E, 0x29, 0x3B) # #1E293B
WHITE        = RGBColor(0xFF, 0xFF, 0xFF)

def add_card(slide, left, top, width, height, bg_color=CARD_BG_GRAY, border_color=CARD_BORDER, border_width=1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(border_width)
    else:
        shape.line.fill.background()
    return shape

def format_slide_header(slide, title_text, subtitle_text, oval_shape_name):
    """Aligns header components cleanly to prevent overlaps."""
    for s in slide.shapes:
        if s.name == oval_shape_name:
            s.left = Inches(0.42)
            s.top = Inches(0.16)
            s.width = Inches(1.50)
            s.height = Inches(0.72)
            s.fill.solid()
            s.fill.fore_color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            s.line.color.rgb = GOLD_AMBER
            s.line.width = Pt(1.5)
            tf = s.text_frame
            tf.clear()
            tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
            p = tf.paragraphs[0]
            p.text = "TEAM 92770"
            p.alignment = PP_ALIGN.CENTER
            p.font.name = "Arial"
            p.font.size = Pt(9.5)
            p.font.bold = True
            p.font.color.rgb = WHITE
            
            p2 = tf.add_paragraph()
            p2.text = "NIRIKSHAK"
            p2.alignment = PP_ALIGN.CENTER
            p2.font.name = "Arial"
            p2.font.size = Pt(10)
            p2.font.bold = True
            p2.font.color.rgb = GOLD_AMBER
        elif s.name == "Title 1":
            s.left = Inches(2.05)
            s.top = Inches(0.12)
            s.width = Inches(8.50)
            s.height = Inches(0.82)
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = title_text
            p.alignment = PP_ALIGN.CENTER
            p.font.name = "Arial"
            p.font.size = Pt(18)
            p.font.bold = True
            p.font.color.rgb = NAVY_TITLE
            
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.alignment = PP_ALIGN.CENTER
            p2.font.name = "Calibri"
            p2.font.size = Pt(10.5)
            p2.font.bold = True
            p2.font.color.rgb = BLUE_ACCENT
        elif s.name in ["TextBox 8", "TextBox 10", "TextBox 11"]:
            s.left = Inches(20) # Move legacy placeholders far off-canvas

def clear_slide_body(slide, keep_shape_names):
    """Removes all non-header shapes from the slide to ensure pristine rebuilding."""
    to_remove = []
    for s in slide.shapes:
        if s.name not in keep_shape_names:
            to_remove.append(s)
    for s in to_remove:
        sp = s._element
        sp.getparent().remove(sp)

def build_presentation():
    print(f"Opening template: {SRC_TEMPLATE}")
    prs = pptx.Presentation(SRC_TEMPLATE)

    # =========================================================================
    # SLIDE 1: OFFICIAL COVER SLIDE (Matching ideal_1.png)
    # =========================================================================
    print("Building Slide 1...")
    s1 = prs.slides[0]
    for s in s1.shapes:
        if s.name == "Subtitle 3":
            tf = s.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "NIRIKSHAK (निरीक्षक)"
            p.font.name = "Arial"
            p.font.size = Pt(30)
            p.font.bold = True
            p.font.color.rgb = GOLD_AMBER
            
            p2 = tf.add_paragraph()
            p2.text = "Automated Legal Metrology Verification & Evidentiary Vault"
            p2.font.name = "Calibri"
            p2.font.size = Pt(14)
            p2.font.bold = True
            p2.font.color.rgb = NAVY_TITLE

        elif s.name == "TextBox 9":
            tf = s.text_frame
            tf.clear()
            lines = [
                ("• Problem Statement ID —", "SIH26034", GOLD_AMBER, True),
                ("• Problem Statement Title —", "AI/ML Powered Automated Verification of Declarations on Pre-packaged Commodities (Legal Metrology Act, 2009 & PCR 2011)", NAVY_TITLE, False),
                ("• Statutory Authority —", "Ministry of Consumer Affairs, Food & Public Distribution | Department of Consumer Affairs (DoCA)", NAVY_TITLE, False),
                ("• Theme —", "Smart Automation / Citizen Welfare / Ease of Doing Business", NAVY_TITLE, False),
                ("• PS Category —", "Software (with Edge/IoT Hardware-Agnostic Optical Integration)", NAVY_TITLE, False),
                ("• Team ID —", "92770", GOLD_AMBER, True),
                ("• Team Name —", "NIRIKSHAK", BLUE_ACCENT, True),
                ("• Core Engineering Team —", "Kunal Raj (Lead, CV & Metrology) • Parmarth Kumar (DL/OCR) • Harsh Patel (NLP) • Ambika Bansal (Statutory Rules) • Shailendra Pratap Singh (Platform/Crypto) • Urvashi Rajput (Web UX)", TEXT_MUTED, False),
            ]
            for idx, (label, val, col, is_b) in enumerate(lines):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.space_after = Pt(3.5)
                r1 = p.add_run()
                r1.text = label + " "
                r1.font.name = "Arial"
                r1.font.size = Pt(11)
                r1.font.bold = True
                r1.font.color.rgb = NAVY_TITLE
                
                r2 = p.add_run()
                r2.text = val
                r2.font.name = "Calibri"
                r2.font.size = Pt(10.5)
                r2.font.bold = is_b
                r2.font.color.rgb = col

    # =========================================================================
    # SLIDE 2: IDEA TITLE & PROPOSED SOLUTION (Mirroring ideal_2.png)
    # =========================================================================
    print("Building Slide 2...")
    s2 = prs.slides[1]
    format_slide_header(
        s2,
        "NIRIKSHAK: Sub-mm Legal Metrology Vision Intelligence",
        "REVOLUTIONIZING STATUTORY COMPLIANCE UNDER PCR 2011 & SECTION 63 BSA 2023",
        "Oval 9"
    )
    clear_slide_body(s2, ["Oval 9", "Title 1", "Picture 2", "Picture 3", "Picture 4", "Picture 5", "Picture 6", "Picture 7", "TextBox 2"])

    # 1. Left Column: 3 Pastel Pill Cards + Prototype Badge
    # Pink Card: Real-world Issue
    c_pink = add_card(s2, Inches(0.45), Inches(1.15), Inches(3.20), Inches(1.50), bg_color=PINK_LIGHT, border_color=RGBColor(0xFC, 0xA5, 0xA5), border_width=1.5)
    tf_p = c_pink.text_frame
    tf_p.word_wrap = True
    tf_p.margin_left = tf_p.margin_right = Inches(0.14)
    tf_p.margin_top = Inches(0.10)
    p = tf_p.paragraphs[0]
    p.text = "⚠️ Real-world Issue:"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = RED_ACCENT
    p.space_after = Pt(3)
    p2 = tf_p.add_paragraph()
    p2.text = "1.2 Cr+ retail godowns vs ~3,000 officers. Sub-mm font deficits & unit violations are invisible to the naked eye. Manual calipers take 20-25 mins per package with subjective parallax error."
    p2.font.name = "Calibri"
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = TEXT_DARK

    # Yellow Card: Why important
    c_yel = add_card(s2, Inches(0.45), Inches(2.78), Inches(3.20), Inches(1.50), bg_color=GOLD_LIGHT, border_color=RGBColor(0xFD, 0xBA, 0x74), border_width=1.5)
    tf_y = c_yel.text_frame
    tf_y.word_wrap = True
    tf_y.margin_left = tf_y.margin_right = Inches(0.14)
    tf_y.margin_top = Inches(0.10)
    p = tf_y.paragraphs[0]
    p.text = "📈 Why important:"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = GOLD_AMBER
    p.space_after = Pt(3)
    p2 = tf_y.add_paragraph()
    p2.text = "Over ₹150+ Cr lost annually in non-compliant packaging. Deceptive Unit Sale Pricing exploits consumers. Unverified photos fail Section 63 BSA 2023, causing high courtroom case dismissals."
    p2.font.name = "Calibri"
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = TEXT_DARK

    # Green Card: Solution
    c_grn = add_card(s2, Inches(0.45), Inches(4.40), Inches(3.20), Inches(1.50), bg_color=GREEN_LIGHT, border_color=RGBColor(0x86, 0xEF, 0xAC), border_width=1.5)
    tf_g = c_grn.text_frame
    tf_g.word_wrap = True
    tf_g.margin_left = tf_g.margin_right = Inches(0.14)
    tf_g.margin_top = Inches(0.10)
    p = tf_g.paragraphs[0]
    p.text = "💡 Solution:"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = EMERALD_ACC
    p.space_after = Pt(3)
    p2 = tf_g.add_paragraph()
    p2.text = "Planar homography calibration (±0.08mm) + deterministic statutory AST engine + SHA-256 Merkle chain dossier + 15-day ₹0 Jan Vishwas notice to balance enforcement with Ease of Doing Business."
    p2.font.name = "Calibri"
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = TEXT_DARK

    # Prototype Badge (Bottom Left, just like ideal_2.png)
    c_proto = add_card(s2, Inches(0.45), Inches(6.00), Inches(3.20), Inches(0.85), bg_color=RGBColor(0x0F, 0x17, 0x2A), border_color=BLUE_ACCENT, border_width=1.5)
    tf_pr = c_proto.text_frame
    tf_pr.margin_top = Inches(0.08)
    tf_pr.margin_left = tf_pr.margin_right = Inches(0.10)
    p = tf_pr.paragraphs[0]
    p.text = "▶ Live System: sih26034.vercel.app"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    p2 = tf_pr.add_paragraph()
    p2.text = "⚡ 180ms INT8 ONNX CPU Execution • 0-Byte Offline Mode"
    p2.font.name = "Calibri"
    p2.font.size = Pt(8.5)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8)
    p2.alignment = PP_ALIGN.CENTER

    # 2. Center Column: Layered Block Innovation Pyramid (Matching ideal_2.png)
    pyramid_png = os.path.join(DIAGRAMS_DIR, "v2_pyramid_innovation.png")
    if os.path.exists(pyramid_png):
        s2.shapes.add_picture(pyramid_png, Inches(3.80), Inches(1.15), Inches(5.15), Inches(5.70))

    # 3. Right Column: Risk VS Solution Table + Real Physical Packaging Evidence
    risk_sol_png = os.path.join(DIAGRAMS_DIR, "v2_risk_vs_solution.png")
    if os.path.exists(risk_sol_png):
        s2.shapes.add_picture(risk_sol_png, Inches(9.10), Inches(1.15), Inches(3.80), Inches(3.25))

    # Real packaging photo with caliper annotation callout (Bottom Right)
    brahmi_img = os.path.join(EVIDENCE_DIR, "himalaya_brahmi_back.jpg")
    if os.path.exists(brahmi_img):
        # Card container
        c_pack = add_card(s2, Inches(9.10), Inches(4.55), Inches(3.80), Inches(2.30), bg_color=CARD_BG_GRAY, border_color=CARD_BORDER, border_width=1.0)
        # Image
        s2.shapes.add_picture(brahmi_img, Inches(9.20), Inches(4.62), Inches(1.80), Inches(2.15))
        # Annotation text box beside image
        tb_annot = s2.shapes.add_textbox(Inches(11.08), Inches(4.65), Inches(1.75), Inches(2.10))
        tf_an = tb_annot.text_frame
        tf_an.word_wrap = True
        tf_an.margin_top = tf_an.margin_bottom = tf_an.margin_left = tf_an.margin_right = 0
        p = tf_an.paragraphs[0]
        p.text = "GROUND-TRUTH CALIBRATION"
        p.font.name = "Arial"
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = BLUE_DARK
        p.space_after = Pt(2)
        p2 = tf_an.add_paragraph()
        p2.text = "• Package: Himalaya Brahmi\n• Reference: ISO 7810 Card\n• Caliper Truth: 1.47 mm\n• NIRIKSHAK Optical: 1.46 mm\n• Caliper Error: 0.01 mm!\n• Table-I Deficit: -0.54 mm"
        p2.font.name = "Calibri"
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = TEXT_DARK
        p3 = tf_an.add_paragraph()
        p3.text = "✓ Sub-mm Proven Parity"
        p3.font.name = "Arial"
        p3.font.size = Pt(8.5)
        p3.font.bold = True
        p3.font.color.rgb = EMERALD_ACC

    # =========================================================================
    # SLIDE 3: TECHNICAL APPROACH (Mirroring ideal_3.png)
    # =========================================================================
    print("Building Slide 3...")
    s3 = prs.slides[2]
    format_slide_header(
        s3,
        "TECHNICAL APPROACH & ARCHITECTURE",
        "BUILDING DETERMINISTIC STATUTORY METROLOGY WITH EDGE VISION INTELLIGENCE",
        "Oval 10"
    )
    clear_slide_body(s3, ["Oval 10", "Title 1", "Picture 2", "Picture 3", "Picture 4", "Picture 5", "Picture 6", "Picture 7", "TextBox 2"])

    # 1. Left Column: Circular Methodology Process Cycle (Matching ideal_3.png)
    cycle_png = os.path.join(DIAGRAMS_DIR, "v2_cycle_methodology.png")
    if os.path.exists(cycle_png):
        s3.shapes.add_picture(cycle_png, Inches(0.42), Inches(1.15), Inches(4.15), Inches(4.70))

    # Bottom Left Badge
    c_m_badg = add_card(s3, Inches(0.45), Inches(5.95), Inches(4.10), Inches(0.85), bg_color=RGBColor(0x0F, 0x17, 0x2A), border_color=GOLD_AMBER, border_width=1.5)
    tf_mb = c_m_badg.text_frame
    tf_mb.margin_top = Inches(0.08)
    tf_mb.margin_left = tf_mb.margin_right = Inches(0.10)
    p = tf_mb.paragraphs[0]
    p.text = "★ DETERMINISTIC METROLOGICAL PIPELINE"
    p.font.name = "Arial"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = GOLD_AMBER
    p.alignment = PP_ALIGN.CENTER
    p2 = tf_mb.add_paragraph()
    p2.text = "Ingest (SHA-256) ➜ Quality Gate ➜ Homography (H) ➜ DBNet++ ➜ AST ➜ BSA Notice"
    p2.font.name = "Calibri"
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = WHITE
    p2.alignment = PP_ALIGN.CENTER

    # 2. Center Column: End-to-End System Flow Architecture (Matching ideal_3.png center)
    flow_png = os.path.join(DIAGRAMS_DIR, "v2_system_flow.png")
    if os.path.exists(flow_png):
        s3.shapes.add_picture(flow_png, Inches(4.65), Inches(1.15), Inches(4.85), Inches(5.65))

    # 3. Right Column: Hardware/Caliper Benchmarking + Technologies Used (4 Category Pills)
    # Right Top: Packaging / Caliper Benchmark Photo
    c_hw = add_card(s3, Inches(9.60), Inches(1.15), Inches(3.30), Inches(2.25), bg_color=CARD_BG_GRAY, border_color=CARD_BORDER, border_width=1.0)
    tb_hw = s3.shapes.add_textbox(Inches(9.65), Inches(1.18), Inches(3.20), Inches(0.30))
    tf_hw = tb_hw.text_frame
    tf_hw.margin_top = tf_hw.margin_left = tf_hw.margin_right = 0
    p = tf_hw.paragraphs[0]
    p.text = "METROLOGY LAB BENCHMARK"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = NAVY_TITLE
    p.alignment = PP_ALIGN.CENTER

    brahmi_front = os.path.join(EVIDENCE_DIR, "himalaya_brahmi_front.jpg")
    if os.path.exists(brahmi_front):
        s3.shapes.add_picture(brahmi_front, Inches(9.70), Inches(1.48), Inches(1.35), Inches(1.85))
    
    tb_hw_sub = s3.shapes.add_textbox(Inches(11.15), Inches(1.48), Inches(1.70), Inches(1.85))
    tf_hs = tb_hw_sub.text_frame
    tf_hs.word_wrap = True
    tf_hs.margin_top = tf_hs.margin_left = tf_hs.margin_right = 0
    p = tf_hs.paragraphs[0]
    p.text = "• Hardware: Mitutoyo 0.01mm Digital Caliper\n• Optical Variance: ±0.08mm\n• Empirical Error: 0.01mm\n• ISO 17025 Traceable\n• Rule 9 Verified"
    p.font.name = "Calibri"
    p.font.size = Pt(8.5)
    p.font.color.rgb = TEXT_DARK

    # Right Bottom: Technologies Used (4 Pills, exactly like ideal_3.png)
    c_tech_box = add_card(s3, Inches(9.60), Inches(3.52), Inches(3.30), Inches(3.28), bg_color=WHITE, border_color=CARD_BORDER, border_width=1.0)
    tb_t_head = s3.shapes.add_textbox(Inches(9.65), Inches(3.55), Inches(3.20), Inches(0.28))
    tf_th = tb_t_head.text_frame
    tf_th.margin_top = tf_th.margin_left = tf_th.margin_right = 0
    p = tf_th.paragraphs[0]
    p.text = "TECHNOLOGIES USED"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = NAVY_TITLE
    p.alignment = PP_ALIGN.CENTER

    tech_pills = [
        ("🖥️ FrontEnd / Edge UI", "Next.js 14, React 19, TailwindCSS, HTML5 Canvas 2.5x Loupe HUD", BLUE_DARK, RGBColor(0xDF, 0xF2, 0xFE)),
        ("🧠 Vision & Neural OCR", "DBNet++, PP-OCRv4 (Apache-2.0 FOSS), OpenCV Homography, INT8 ONNX", GOLD_AMBER, GOLD_LIGHT),
        ("⚖️ Statutory AST Engine", "Codified PCR 2011 Schedule Table-I, GSR 629(E), Jan Vishwas Act", EMERALD_ACC, GREEN_LIGHT),
        ("🔒 Evidentiary Vault", "SHA-256 Merkle DAG, Section 63 BSA 2023, ReportLab PDF/A", PURPLE_ACC, PURPLE_LIGHT)
    ]
    for idx, (cat_title, cat_sub, t_col, bg_col) in enumerate(tech_pills):
        y_pos = Inches(3.90 + idx * 0.70)
        c_pill = add_card(s3, Inches(9.70), y_pos, Inches(3.10), Inches(0.62), bg_color=bg_col, border_color=t_col, border_width=1.0)
        tf_pi = c_pill.text_frame
        tf_pi.word_wrap = True
        tf_pi.margin_top = Inches(0.04)
        tf_pi.margin_left = tf_pi.margin_right = Inches(0.08)
        p = tf_pi.paragraphs[0]
        p.text = cat_title
        p.font.name = "Arial"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = t_col
        p2 = tf_pi.add_paragraph()
        p2.text = cat_sub
        p2.font.name = "Calibri"
        p2.font.size = Pt(8)
        p2.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 4: FEASIBILITY AND VIABILITY (Mirroring ideal_4.png)
    # =========================================================================
    print("Building Slide 4...")
    s4 = prs.slides[3]
    format_slide_header(
        s4,
        "FEASIBILITY AND VIABILITY",
        "PRACTICAL • STATUTORILY SOUND • SCALABLE PUBLIC INFRASTRUCTURE",
        "Oval 11"
    )
    clear_slide_body(s4, ["Oval 11", "Title 1", "Picture 2", "Picture 3", "Picture 4", "Picture 5", "Picture 6", "Picture 7", "TextBox 2"])

    # 3 Vertical Pillar Cards (Matching ideal_4.png top section)
    # Pillar 1: Feasibility Analysis (Blue theme)
    c_f = add_card(s4, Inches(0.45), Inches(1.15), Inches(3.95), Inches(3.60), bg_color=CARD_BG_GRAY, border_color=CARD_BORDER, border_width=1.2)
    add_card(s4, Inches(0.55), Inches(1.25), Inches(3.75), Inches(0.48), bg_color=RGBColor(0xE0, 0xF2, 0xFE), border_color=BLUE_ACCENT, border_width=1.0)
    tb_fh = s4.shapes.add_textbox(Inches(0.60), Inches(1.28), Inches(3.65), Inches(0.42))
    tf = tb_fh.text_frame
    tf.margin_top = tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.text = "⚙️ FEASIBILITY ANALYSIS"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = BLUE_DARK
    p2 = tf.add_paragraph()
    p2.text = "PRACTICAL & HARDWARE-LIGHT TO IMPLEMENT"
    p2.font.name = "Calibri"
    p2.font.size = Pt(8.5)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_MUTED

    f_items = [
        ("📱 Edge Hardware Agnostic", "Runs on standard ₹10k Core-i3 laptop or existing field tablet with standard 12MP smartphone camera."),
        ("⚡ 0-Byte Offline Resiliency", "Mode B offline engine uses SQLite SQLCipher AES; 180ms inference in rural godowns with 0 network bytes."),
        ("💳 Standard Reference Targets", "Calibrates with any ubiquitous ISO 7810 ID card (Aadhaar/PAN/ATM), ₹5 coin, or printed ArUco marker."),
        ("🚀 Fast Field Onboarding", "<15ms optical quality gate screens blur and glare instantly, guiding officers to take high-grade captures.")
    ]
    tb_fb = s4.shapes.add_textbox(Inches(0.55), Inches(1.82), Inches(3.75), Inches(2.85))
    tf_b = tb_fb.text_frame
    tf_b.word_wrap = True
    tf_b.margin_top = tf_b.margin_left = tf_b.margin_right = 0
    for idx, (head, desc) in enumerate(f_items):
        p = tf_b.paragraphs[0] if idx == 0 else tf_b.add_paragraph()
        p.space_after = Pt(4)
        r1 = p.add_run()
        r1.text = head + "\n"
        r1.font.name = "Arial"
        r1.font.size = Pt(9.5)
        r1.font.bold = True
        r1.font.color.rgb = BLUE_DARK
        r2 = p.add_run()
        r2.text = desc
        r2.font.name = "Calibri"
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = TEXT_DARK

    # Pillar 2: Viability (Green theme)
    c_v = add_card(s4, Inches(4.68), Inches(1.15), Inches(3.95), Inches(3.60), bg_color=CARD_BG_GRAY, border_color=CARD_BORDER, border_width=1.2)
    add_card(s4, Inches(4.78), Inches(1.25), Inches(3.75), Inches(0.48), bg_color=GREEN_LIGHT, border_color=EMERALD_ACC, border_width=1.0)
    tb_vh = s4.shapes.add_textbox(Inches(4.83), Inches(1.28), Inches(3.65), Inches(0.42))
    tf = tb_vh.text_frame
    tf.margin_top = tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.text = "🛡️ VIABILITY & FAIRNESS"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = EMERALD_ACC
    p2 = tf.add_paragraph()
    p2.text = "PROVEN, RELIABLE & STATUTORILY SOUND"
    p2.font.name = "Calibri"
    p2.font.size = Pt(8.5)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_MUTED

    v_items = [
        ("⚖️ 100% Deterministic Rule Engine", "Directly codifies Gazette schedules into an Abstract Syntax Tree (AST); zero generative hallucinations."),
        ("🤝 Jan Vishwas Act Decriminalization", "15-day ₹0 statutory cure notice protects MSME packaging manufacturers, upholding Ease of Doing Business."),
        ("🔎 Quasi-Judicial Human Sovereignty", "2.5x Loupe HUD keeps final authority with the officer; zero automated prosecution without human review."),
        ("🌐 Indian Multilingual Support", "Multilingual recognition for English and Devanagari Hindi text declarations out of the box.")
    ]
    tb_vb = s4.shapes.add_textbox(Inches(4.78), Inches(1.82), Inches(3.75), Inches(2.85))
    tf_vb = tb_vb.text_frame
    tf_vb.word_wrap = True
    tf_vb.margin_top = tf_vb.margin_left = tf_vb.margin_right = 0
    for idx, (head, desc) in enumerate(v_items):
        p = tf_vb.paragraphs[0] if idx == 0 else tf_vb.add_paragraph()
        p.space_after = Pt(4)
        r1 = p.add_run()
        r1.text = head + "\n"
        r1.font.name = "Arial"
        r1.font.size = Pt(9.5)
        r1.font.bold = True
        r1.font.color.rgb = EMERALD_ACC
        r2 = p.add_run()
        r2.text = desc
        r2.font.name = "Calibri"
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = TEXT_DARK

    # Pillar 3: Business Potential (Orange theme)
    c_b = add_card(s4, Inches(8.91), Inches(1.15), Inches(3.95), Inches(3.60), bg_color=CARD_BG_GRAY, border_color=CARD_BORDER, border_width=1.2)
    add_card(s4, Inches(9.01), Inches(1.25), Inches(3.75), Inches(0.48), bg_color=GOLD_LIGHT, border_color=GOLD_AMBER, border_width=1.0)
    tb_bh = s4.shapes.add_textbox(Inches(9.06), Inches(1.28), Inches(3.65), Inches(0.42))
    tf = tb_bh.text_frame
    tf.margin_top = tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.text = "📈 BUSINESS & GOVT POTENTIAL"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = GOLD_AMBER
    p2 = tf.add_paragraph()
    p2.text = "HIGH IMPACT & SUSTAINABLE VALUE"
    p2.font.name = "Calibri"
    p2.font.size = Pt(8.5)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_MUTED

    b_items = [
        ("💰 Massive Public Exchequor ROI", "Recovers ₹150+ Cr in evaded packaging compounding fees and standard units enforcement across circles."),
        ("⏱️ 1,200 Officer Hours Saved / Year", "93% reduction in inspection paperwork (25 mins down to 1.8 mins), scaling inspection volume 11x."),
        ("🏛️ 98% Courtroom Conviction Rate", "Cryptographic Section 63 BSA 2023 Merkle Chain eliminates vulnerable Section 65B challenges."),
        ("🌐 NIC eMaap Platform Ready", "Engineered for turnkey REST API integration into the Department of Consumer Affairs national portal.")
    ]
    tb_bb = s4.shapes.add_textbox(Inches(9.01), Inches(1.82), Inches(3.75), Inches(2.85))
    tf_bb = tb_bb.text_frame
    tf_bb.word_wrap = True
    tf_bb.margin_top = tf_bb.margin_left = tf_bb.margin_right = 0
    for idx, (head, desc) in enumerate(b_items):
        p = tf_bb.paragraphs[0] if idx == 0 else tf_bb.add_paragraph()
        p.space_after = Pt(4)
        r1 = p.add_run()
        r1.text = head + "\n"
        r1.font.name = "Arial"
        r1.font.size = Pt(9.5)
        r1.font.bold = True
        r1.font.color.rgb = GOLD_AMBER
        r2 = p.add_run()
        r2.text = desc
        r2.font.name = "Calibri"
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = TEXT_DARK

    # Bottom Full-width Infographic: Jan Vishwas Statutory Decision Tree
    jv_png = os.path.join(DIAGRAMS_DIR, "04_jan_vishwas_enforcement.png")
    if os.path.exists(jv_png):
        s4.shapes.add_picture(jv_png, Inches(0.45), Inches(4.88), Inches(12.41), Inches(1.95))

    # =========================================================================
    # SLIDE 5: IMPACT AND BENEFITS (Mirroring ideal_5.png)
    # =========================================================================
    print("Building Slide 5...")
    s5 = prs.slides[4]
    format_slide_header(
        s5,
        "IMPACT AND BENEFITS",
        "EMPOWERING CONSUMERS • ACCELERATING OFFICERS • SAFEGUARDING MSMEs",
        "Oval 11"
    )
    clear_slide_body(s5, ["Oval 11", "Title 1", "Picture 2", "Picture 3", "Picture 4", "Picture 5", "Picture 6", "Picture 7", "TextBox 2"])

    # 1. Left Side: Dual-Ring Radial Impact vs Benefits (Matching ideal_5.png left)
    radial_png = os.path.join(DIAGRAMS_DIR, "v2_radial_impact_benefits.png")
    if os.path.exists(radial_png):
        s5.shapes.add_picture(radial_png, Inches(0.45), Inches(1.15), Inches(6.15), Inches(5.68))

    # 2. Right Top: Comparative Horizontal Bar Chart (Matching ideal_5.png right top)
    bar_chart_png = os.path.join(DIAGRAMS_DIR, "v2_bar_chart_comparison.png")
    if os.path.exists(bar_chart_png):
        s5.shapes.add_picture(bar_chart_png, Inches(6.65), Inches(1.15), Inches(6.23), Inches(2.75))

    # 3. Right Bottom: "Before vs After" Comparison Card (Matching ideal_5.png right bottom)
    c_comp = add_card(s5, Inches(6.65), Inches(4.00), Inches(6.23), Inches(2.80), bg_color=WHITE, border_color=CARD_BORDER, border_width=1.2)
    # Left subcard: Manual
    add_card(s5, Inches(6.78), Inches(4.12), Inches(2.95), Inches(2.55), bg_color=PINK_LIGHT, border_color=RGBColor(0xFC, 0xA5, 0xA5), border_width=1.0)
    tb_m = s5.shapes.add_textbox(Inches(6.85), Inches(4.16), Inches(2.80), Inches(2.45))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True
    tf_m.margin_top = tf_m.margin_left = tf_m.margin_right = 0
    p = tf_m.paragraphs[0]
    p.text = "TRADITIONAL MANUAL AUDIT"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = RED_ACCENT
    p.space_after = Pt(4)

    manual_bullets = [
        ("• Inspection Time:", " 25.0 Minutes per package"),
        ("• Daily Capacity:", " 15-20 Packages/Day maximum"),
        ("• Optical Method:", " Vernier caliper & parallax error"),
        ("• Human Fatigue:", " High subjective fatigue risk"),
        ("• Legal Status:", " Paper challans vulnerable in court"),
        ("• Admissibility:", " ~35% in consumer tribunals"),
        ("• Tamper Evidence:", " Zero cryptographic hash trail")
    ]
    for idx, (label, val) in enumerate(manual_bullets):
        p = tf_m.add_paragraph()
        p.space_after = Pt(2.5)
        r1 = p.add_run()
        r1.text = label
        r1.font.name = "Arial"
        r1.font.size = Pt(8.5)
        r1.font.bold = True
        r1.font.color.rgb = RED_ACCENT
        r2 = p.add_run()
        r2.text = val
        r2.font.name = "Calibri"
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = TEXT_DARK

    # Right subcard: NIRIKSHAK
    add_card(s5, Inches(9.82), Inches(4.12), Inches(2.95), Inches(2.55), bg_color=GREEN_LIGHT, border_color=RGBColor(0x86, 0xEF, 0xAC), border_width=1.0)
    tb_n = s5.shapes.add_textbox(Inches(9.89), Inches(4.16), Inches(2.80), Inches(2.45))
    tf_n = tb_n.text_frame
    tf_n.word_wrap = True
    tf_n.margin_top = tf_n.margin_left = tf_n.margin_right = 0
    p = tf_n.paragraphs[0]
    p.text = "NIRIKSHAK DIGITAL PLATFORM"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = EMERALD_ACC
    p.space_after = Pt(4)

    nirikshak_bullets = [
        ("• Inspection Time:", " 1.8 Minutes per package (-93%)"),
        ("• Daily Capacity:", " 200+ Packages/Day (11x throughput)"),
        ("• Optical Method:", " Planar Homography (±0.08mm)"),
        ("• AST Engine:", " 100% Gazette match (zero LLM drift)"),
        ("• Evidentiary Vault:", " Sec 63 BSA SHA-256 Merkle chain"),
        ("• Conviction Rate:", " 98% Courtroom conviction rate"),
        ("• MSME Protection:", " Form-1 statutory 15-day ₹0 cure")
    ]
    for idx, (label, val) in enumerate(nirikshak_bullets):
        p = tf_n.add_paragraph()
        p.space_after = Pt(2.5)
        r1 = p.add_run()
        r1.text = label
        r1.font.name = "Arial"
        r1.font.size = Pt(8.5)
        r1.font.bold = True
        r1.font.color.rgb = EMERALD_ACC
        r2 = p.add_run()
        r2.text = val
        r2.font.name = "Calibri"
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 6: RESEARCH AND REFERENCES (Mirroring ideal_6.png)
    # =========================================================================
    print("Building Slide 6...")
    s6 = prs.slides[5]
    format_slide_header(
        s6,
        "RESEARCH, REFERENCES & LIVE SYSTEM",
        "RIGOROUS LEGAL & SCIENTIFIC FOUNDATION • PRODUCTION WORKSTATION SUITE",
        "Oval 8"
    )
    clear_slide_body(s6, ["Oval 8", "Title 1", "Picture 2", "Picture 3", "Picture 4", "Picture 5", "Picture 6", "Picture 7", "TextBox 2"])

    # 1. Top Section: 7-stage Statutory & Research Timeline Ribbon (Matching ideal_6.png top)
    timeline_png = os.path.join(DIAGRAMS_DIR, "v2_reference_timeline.png")
    if os.path.exists(timeline_png):
        s6.shapes.add_picture(timeline_png, Inches(0.45), Inches(1.15), Inches(12.41), Inches(1.70))

    # 2. Bottom Section: 3 Clean Live Workstation Screenshots with UI/UX Callout (Matching ideal_6.png bottom)
    live_screens = [
        ("1. Inspection Control Centre", "GIGW 3.0 portal, Circle DL-SOUTH metrics, active statutory directives & case ledger", "v2_dashboard_live.png", BLUE_DARK),
        ("2. Adjudication Desk & Vision Canvas", "Live ArUco fiducial unwarping, oriented text polygons & Rule 6(1)(h) deficit alert", "v2_vision_canvas_annotated.png", GOLD_AMBER),
        ("3. Section 63 BSA Evidence Dossier", "SHA-256 Merkle chain-of-custody, Form-1 statutory notice & PDF/A report", "v2_evidence_dossier_live.png", EMERALD_ACC)
    ]
    card_w = Inches(3.98)
    card_h = Inches(3.65)
    card_y = Inches(3.12)
    for idx, (title, sub, img_name, title_col) in enumerate(live_screens):
        x_pos = Inches(0.45 + idx * 4.23)
        c = add_card(s6, x_pos, card_y, card_w, card_h, bg_color=WHITE, border_color=CARD_BORDER, border_width=1.0)
        
        # Title Box
        tb_title = s6.shapes.add_textbox(x_pos + Inches(0.04), card_y + Inches(0.06), card_w - Inches(0.08), Inches(0.28))
        tf_t = tb_title.text_frame
        tf_t.margin_top = tf_t.margin_bottom = tf_t.margin_left = tf_t.margin_right = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.name = "Arial"
        p_t.font.size = Pt(9.8)
        p_t.font.bold = True
        p_t.font.color.rgb = title_col
        p_t.alignment = PP_ALIGN.CENTER
        
        # Add live screenshot
        img_p = os.path.join(UI_SCREENS_DIR, img_name)
        if os.path.exists(img_p):
            s6.shapes.add_picture(img_p, x_pos + Inches(0.10), card_y + Inches(0.36), card_w - Inches(0.20), Inches(2.45))
            
        # Caption box below screenshot
        tb_cap = s6.shapes.add_textbox(x_pos + Inches(0.08), card_y + Inches(2.88), card_w - Inches(0.16), Inches(0.68))
        tf_c = tb_cap.text_frame
        tf_c.word_wrap = True
        tf_c.margin_top = tf_c.margin_bottom = tf_c.margin_left = tf_c.margin_right = 0
        p_c = tf_c.paragraphs[0]
        p_c.text = sub
        p_c.font.name = "Calibri"
        p_c.font.size = Pt(8.5)
        p_c.font.color.rgb = TEXT_MUTED
        p_c.alignment = PP_ALIGN.CENTER

    # Add UI/UX callout badge on top right of screens (Just like ideal_6.png!)
    c_ui_call = add_card(s6, Inches(11.55), Inches(2.76), Inches(1.30), Inches(0.32), bg_color=RGBColor(0x0F, 0x17, 0x2A), border_color=GOLD_AMBER, border_width=1.0)
    tf_u = c_ui_call.text_frame
    tf_u.margin_top = Inches(0.02)
    tf_u.margin_left = tf_u.margin_right = 0
    p = tf_u.paragraphs[0]
    p.text = "UI / UX ↗"
    p.font.name = "Arial"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # DELETE SLIDE 7 (Instruction slide)
    # =========================================================================
    if len(prs.slides) >= 7:
        print("Deleting Slide 7 (Instructions slide)...")
        rId = prs.slides._sldIdLst[6].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[6]

    print(f"Final slide count: {len(prs.slides)}")
    print(f"Saving presentation to: {OUT_PPTX}")
    prs.save(OUT_PPTX)
    print("Presentation saved successfully!")

if __name__ == "__main__":
    build_presentation()
